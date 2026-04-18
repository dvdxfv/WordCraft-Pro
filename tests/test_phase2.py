"""
Phase 2 单元测试

测试所有文件解析器 + 真实文件验证。
"""

import json
import os
import sys
import tempfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.document_model import ElementType, Alignment
from parsers.base import BaseParser
from parsers.txt_parser import TxtParser
from parsers.md_parser import MdParser
from parsers.xlsx_parser import XlsxParser
from parsers.pdf_parser import PdfParser
from parsers.docx_parser import DocxParser
from parsers.dispatcher import parse_file, get_supported_formats, get_parser_for_file


# ============================================================
# 辅助函数：创建临时测试文件
# ============================================================

def _create_temp_file(content: str, suffix: str) -> str:
    """创建临时文件并返回路径"""
    fd, path = tempfile.mkstemp(suffix=suffix)
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(content)
    return path


# ============================================================
# 测试 TXT 解析器
# ============================================================

class TestTxtParser:
    """测试 TXT 解析器"""

    def test_supports_txt(self):
        parser = TxtParser()
        assert parser.supports("test.txt")
        assert not parser.supports("test.docx")

    def test_parse_simple_text(self):
        content = "这是第一段正文内容。\n\n这是第二段正文内容。"
        path = _create_temp_file(content, ".txt")
        try:
            parser = TxtParser()
            doc = parser.parse(path)
            assert doc.source_format == "txt"
            assert len(doc.elements) == 2
            assert doc.elements[0].element_type == ElementType.PARAGRAPH
            assert "第一段" in doc.elements[0].content
        finally:
            os.unlink(path)

    def test_detect_chinese_heading(self):
        content = "一、研究背景\n\n这是正文段落。\n\n二、研究方法\n\n这是方法段落。"
        path = _create_temp_file(content, ".txt")
        try:
            doc = TxtParser().parse(path)
            headings = doc.get_headings()
            assert len(headings) == 2
            assert headings[0].level == 1
            assert "研究背景" in headings[0].content
        finally:
            os.unlink(path)

    def test_detect_reference(self):
        content = "[1] 张三. 论文标题[J]. 期刊名, 2024, 1(1): 1-10.\n\n[2] 李四. 另一篇论文[M]. 出版社, 2023."
        path = _create_temp_file(content, ".txt")
        try:
            doc = TxtParser().parse(path)
            refs = doc.get_references()
            assert len(refs) == 2
        finally:
            os.unlink(path)

    def test_file_not_found(self):
        parser = TxtParser()
        with pytest.raises(FileNotFoundError):
            parser.parse("/nonexistent/file.txt")

    def test_get_all_text(self):
        content = "标题\n\n正文段落一\n\n正文段落二"
        path = _create_temp_file(content, ".txt")
        try:
            doc = TxtParser().parse(path)
            text = doc.get_all_text()
            assert "标题" in text
            assert "正文段落一" in text
        finally:
            os.unlink(path)


# ============================================================
# 测试 Markdown 解析器
# ============================================================

class TestMdParser:
    """测试 Markdown 解析器"""

    def test_supports_md(self):
        parser = MdParser()
        assert parser.supports("test.md")
        assert parser.supports("test.markdown")

    def test_parse_headings(self):
        content = """# 第一章 绪论

## 1.1 研究背景

这是正文内容。

## 1.2 研究目的

### 1.2.1 具体目标

详细描述。
"""
        path = _create_temp_file(content, ".md")
        try:
            doc = MdParser().parse(path)
            headings = doc.get_headings()
            assert len(headings) == 4
            assert headings[0].level == 1
            assert headings[1].level == 2
            assert headings[3].level == 3
        finally:
            os.unlink(path)

    def test_parse_list(self):
        content = """# 列表测试

- 项目一
- 项目二
- 项目三

1. 有序一
2. 有序二
"""
        path = _create_temp_file(content, ".md")
        try:
            doc = MdParser().parse(path)
            lists = doc.find_elements(ElementType.LIST)
            assert len(lists) == 2
            assert len(lists[0].children) == 3
            assert len(lists[1].children) == 2
        finally:
            os.unlink(path)

    def test_parse_table(self):
        content = """# 表格测试

| 姓名 | 年龄 | 城市 |
|------|------|------|
| 张三 | 25 | 北京 |
| 李四 | 30 | 上海 |
"""
        path = _create_temp_file(content, ".md")
        try:
            doc = MdParser().parse(path)
            assert len(doc.tables) == 1
            assert doc.tables[0].num_rows == 3
            assert doc.tables[0].num_cols == 3
            assert doc.tables[0].rows[1][0].content == "张三"
            assert doc.tables[0].rows[0][0].is_header is True
        finally:
            os.unlink(path)

    def test_parse_code_block(self):
        content = """# 代码测试

```python
def hello():
    print("Hello, World!")
```
"""
        path = _create_temp_file(content, ".md")
        try:
            doc = MdParser().parse(path)
            code_blocks = doc.find_elements(ElementType.CODE_BLOCK)
            assert len(code_blocks) == 1
            assert "def hello" in code_blocks[0].content
            assert code_blocks[0].metadata.get("language") == "python"
        finally:
            os.unlink(path)

    def test_clean_markdown_formatting(self):
        content = """# 标题

这是**加粗**和*斜体*的文字，还有`代码`和[链接](https://example.com)。
"""
        path = _create_temp_file(content, ".md")
        try:
            doc = MdParser().parse(path)
            paras = doc.get_paragraphs()
            assert len(paras) == 1
            text = paras[0].content
            assert "**" not in text
            assert "*" not in text
            assert "`" not in text
            assert "加粗" in text
            assert "斜体" in text
            assert "代码" in text
            assert "链接" in text
        finally:
            os.unlink(path)


# ============================================================
# 测试 XLSX 解析器
# ============================================================

class TestXlsxParser:
    """测试 XLSX 解析器"""

    def test_supports_xlsx(self):
        parser = XlsxParser()
        assert parser.supports("test.xlsx")
        assert not parser.supports("test.xls")

    def test_parse_xlsx(self):
        """测试解析真实 Excel 文件"""
        from openpyxl import Workbook
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "测试数据"
            ws.append(["姓名", "年龄", "城市"])
            ws.append(["张三", 25, "北京"])
            ws.append(["李四", 30, "上海"])
            wb.save(path)

            doc = XlsxParser().parse(path)
            assert doc.source_format == "xlsx"
            assert len(doc.tables) == 1
            table = doc.tables[0]
            assert table.num_rows == 3
            assert table.num_cols == 3
            assert table.rows[1][0].content == "张三"
            assert table.rows[0][0].is_header is True
        finally:
            os.close(fd)
            os.unlink(path)


# ============================================================
# 测试 PDF 解析器
# ============================================================

class TestPdfParser:
    """测试 PDF 解析器"""

    def test_supports_pdf(self):
        parser = PdfParser()
        assert parser.supports("test.pdf")
        assert not parser.supports("test.txt")

    def test_parse_pdf(self):
        """测试解析生成的 PDF 文件"""
        # 创建一个简单的 PDF 用于测试
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        fd, path = tempfile.mkstemp(suffix=".pdf")
        try:
            c = canvas.Canvas(path, pagesize=A4)
            # 使用内置字体（避免中文字体问题）
            c.setFont("Helvetica", 16)
            c.drawCentredString(A4[0] / 2, A4[1] - 2 * cm, "Chapter 1 Introduction")
            c.setFont("Helvetica", 12)
            c.drawString(2 * cm, A4[1] - 4 * cm, "This is the first paragraph of content.")
            c.drawString(2 * cm, A4[1] - 5 * cm, "This is the second paragraph.")
            c.showPage()
            c.save()

            doc = PdfParser().parse(path)
            assert doc.source_format == "pdf"
            assert len(doc.elements) > 0
            text = doc.get_all_text()
            assert "Introduction" in text or "introduction" in text.lower()
        except ImportError:
            pytest.skip("reportlab 未安装，跳过 PDF 测试")
        finally:
            os.close(fd)
            if os.path.exists(path):
                os.unlink(path)


# ============================================================
# 测试解析器调度器
# ============================================================

class TestDispatcher:
    """测试解析器调度器"""

    def test_get_supported_formats(self):
        formats = get_supported_formats()
        assert ".txt" in formats
        assert ".md" in formats
        assert ".docx" in formats
        assert ".xlsx" in formats
        assert ".pdf" in formats
        assert ".doc" in formats

    def test_get_parser_for_file(self):
        assert isinstance(get_parser_for_file("test.txt"), TxtParser)
        assert isinstance(get_parser_for_file("test.md"), MdParser)
        assert isinstance(get_parser_for_file("test.xlsx"), XlsxParser)
        assert isinstance(get_parser_for_file("test.pdf"), PdfParser)
        assert isinstance(get_parser_for_file("test.docx"), DocxParser)
        assert isinstance(get_parser_for_file("test.doc"), DocxParser)

    def test_unsupported_format(self):
        with pytest.raises(ValueError, match="不支持的文件格式"):
            parse_file("test.xyz")

    def test_file_not_found(self):
        with pytest.raises(FileNotFoundError):
            parse_file("/nonexistent/file.txt")

    def test_dispatch_txt(self):
        content = "测试内容"
        path = _create_temp_file(content, ".txt")
        try:
            doc = parse_file(path)
            assert doc.source_format == "txt"
        finally:
            os.unlink(path)

    def test_dispatch_md(self):
        content = "# 标题\n\n正文内容"
        path = _create_temp_file(content, ".md")
        try:
            doc = parse_file(path)
            assert doc.source_format == "md"
            assert len(doc.get_headings()) == 1
        finally:
            os.unlink(path)


# ============================================================
# 测试真实文件解析
# ============================================================

class TestRealFiles:
    """使用真实上传的文件进行验证"""

    # 真实文件路径
    DOCX_FILE = "/workspace/.uploads/c0b6893f-cf16-4263-be25-6a16d93da746_附件7：太原市尖草坪区绩效评价报告基本排版要求.docx"

    def test_parse_real_docx(self):
        """测试解析真实的 DOCX 排版要求文件"""
        if not os.path.exists(self.DOCX_FILE):
            pytest.skip("真实 DOCX 文件不存在")

        doc = parse_file(self.DOCX_FILE)

        # 基本验证
        assert doc.source_format == "docx"
        assert len(doc.elements) > 0

        # 应该包含排版要求相关的内容
        text = doc.get_all_text()
        assert len(text) > 50  # 文件应该有内容

        # 验证序列化
        json_str = doc.to_json()
        assert "太原" in json_str or "绩效" in json_str or "排版" in json_str

        # 验证反序列化
        restored = type(doc).from_json(json_str)
        assert len(restored.elements) == len(doc.elements)

        print(f"\n[真实文件解析] DOCX:")
        print(f"  元素数: {len(doc.elements)}")
        print(f"  表格数: {len(doc.tables)}")
        print(f"  节数: {len(doc.sections)}")
        counts = doc.element_count()
        for k, v in counts.items():
            print(f"  {k}: {v}")
        print(f"  前200字: {text[:200]}...")

    def test_parse_real_docx_page_setup(self):
        """测试真实 DOCX 文件的页面设置提取"""
        if not os.path.exists(self.DOCX_FILE):
            pytest.skip("真实 DOCX 文件不存在")

        doc = parse_file(self.DOCX_FILE)

        # 验证页面设置
        assert doc.page_setup.paper_size == "A4"
        assert doc.page_setup.margin_top_cm > 0
        assert doc.page_setup.margin_bottom_cm > 0

        print(f"\n[页面设置]")
        print(f"  纸张: {doc.page_setup.paper_size}")
        print(f"  上边距: {doc.page_setup.margin_top_cm}cm")
        print(f"  下边距: {doc.page_setup.margin_bottom_cm}cm")
        print(f"  左边距: {doc.page_setup.margin_left_cm}cm")
        print(f"  右边距: {doc.page_setup.margin_right_cm}cm")

    def test_parse_real_docx_styles(self):
        """测试真实 DOCX 文件的样式提取"""
        if not os.path.exists(self.DOCX_FILE):
            pytest.skip("真实 DOCX 文件不存在")

        doc = parse_file(self.DOCX_FILE)

        # 检查是否有提取到字体信息
        has_font = False
        for elem in doc.elements:
            if elem.font_style.font_name_cn or elem.font_style.font_name_en:
                has_font = True
                break

        print(f"\n[样式提取]")
        print(f"  提取到字体信息: {'是' if has_font else '否'}")

        # 打印前几个元素的样式信息
        for i, elem in enumerate(doc.elements[:5]):
            fs = elem.font_style
            print(f"  元素{i}: type={elem.element_type.value}, "
                  f"font_cn={fs.font_name_cn or '(空)'}, "
                  f"size={fs.font_size_pt or '(空)'}pt, "
                  f"bold={fs.bold}")


if __name__ == "__main__":
    pytest.main([__file__, "-v", "-s"])
